"""Génération de fichiers PDF et Excel pour les rapports."""
from html import escape
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from rest_framework.exceptions import ValidationError


EXCEL_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_report_response(request, title, filename, rows, columns):
    export_format = request.query_params.get("export", "").strip().lower()
    if not export_format:
        return None
    if export_format not in {"xlsx", "pdf"}:
        raise ValidationError({"export": "Le format doit être « xlsx » ou « pdf »."})

    output = BytesIO()
    if export_format == "xlsx":
        _write_excel(output, title, rows, columns)
        content_type = EXCEL_MIME_TYPE
    else:
        _write_pdf(output, title, rows, columns)
        content_type = "application/pdf"

    response = HttpResponse(output.getvalue(), content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}.{export_format}"'
    return response


def _write_excel(output, title, rows, columns):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = title[:31]
    worksheet.append([title])
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    worksheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="173B75")
    worksheet["A1"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 28
    worksheet.append([label for _, label in columns])
    for cell in worksheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2457D6")

    for row in rows:
        worksheet.append([_excel_safe_value(row.get(key)) for key, _ in columns])

    worksheet.freeze_panes = "A3"
    worksheet.auto_filter.ref = worksheet.dimensions
    for column_index, column in enumerate(worksheet.columns, start=1):
        content_width = max(
            (len(str(cell.value or "")) for cell in column),
            default=10,
        )
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(content_width + 2, 12), 48
        )
        for cell in column[2:]:
            if isinstance(cell.value, (int, float)) or cell.data_type == "n":
                cell.number_format = "#,##0.00"
    workbook.save(output)


def _excel_safe_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value


def _write_pdf(output, title, rows, columns):
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=14 * mm,
        leftMargin=14 * mm,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        title=title,
    )
    styles = getSampleStyleSheet()
    styles["BodyText"].fontSize = 8
    styles["BodyText"].leading = 10
    table_data = [
        [Paragraph(escape(label), styles["BodyText"]) for _, label in columns]
    ]
    table_data.extend(
        [
            [
                Paragraph(escape(str(row.get(key) if row.get(key) is not None else "—")), styles["BodyText"])
                for key, _ in columns
            ]
            for row in rows
        ]
    )
    table = Table(
        table_data,
        colWidths=[document.width / len(columns)] * len(columns),
        repeatRows=1,
        hAlign="LEFT",
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2457D6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D0D7E2")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FB")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    document.build(
        [
            Paragraph(escape(title), styles["Title"]),
            Spacer(1, 6 * mm),
            table,
        ]
    )
